#!/usr/bin/env python3
"""Word note for Week 1–3 inverse kinematics."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def font(run, size=12, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def para(doc, text, **kw):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(kw.get("sa", 6))
    if kw.get("center"):
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    font(r, size=kw.get("size", 12), bold=kw.get("bold", False), italic=kw.get("italic", False))
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    font(r)
    return p


def add_table(doc, rows):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, txt in enumerate(row):
            cell = t.rows[i].cells[j]
            cell.text = ""
            r = cell.paragraphs[0].add_run(str(txt))
            font(r, size=9, bold=(i == 0))
    doc.add_paragraph()


def build():
    doc = Document()
    sec = doc.sections[0]
    for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, m, Inches(1))

    para(doc, "Week 1–3 Deliverable 3", bold=True, center=True, sa=2)
    para(doc, "Inverse Kinematics of the Hanging Planar Five-Bar", size=16, bold=True, center=True, sa=4)
    para(doc, "Bo Zhang (8571260)  |  Dr Chin-Hsing Kuo  |  September 2026", center=True, sa=10)

    para(doc, "1. What was done", bold=True, size=14)
    para(
        doc,
        "A Python inverse-kinematics program was written for the frozen Week 1–3 geometry (d = 180 mm, l1 = l2 = 120 mm, l3 = l4 = 180 mm). Given the end-effector E = (x, y), it returns the crank angles (θ_L, θ_R). Five checkpoint poses were checked by reconstructing the four link lengths and by running forward kinematics. The CAD sample pose is recovered exactly. GSM is still not included.",
    )
    para(doc, "Files: scripts/five_bar_kinematics.py, scripts/verify_five_bar_ik.py, scripts/test_five_bar_ik.py, docs/week1-3/ik_checkpoints.csv.")

    para(doc, "2. Closed-loop equations", bold=True, size=14)
    para(doc, "Origin at the midpoint of the ground. A = (−d/2, 0), B = (d/2, 0). Crank angles from +x, counterclockwise:")
    para(doc, "P = A + l1 [cos θ_L,  sin θ_L]", italic=True)
    para(doc, "Q = B + l2 [cos θ_R,  sin θ_R]", italic=True)
    para(doc, "||E − P|| = l3,     ||E − Q|| = l4", italic=True)
    para(
        doc,
        "These four scalar conditions are the loop closure. Inverse kinematics treats each leg as a two-circle intersection: P is an intersection of circle(A, l1) and circle(E, l3); Q is an intersection of circle(B, l2) and circle(E, l4). Then θ_L = atan2(P_y − A_y, P_x − A_x) and likewise for θ_R.",
    )

    para(doc, "3. How many solutions", bold=True, size=14)
    para(
        doc,
        "Each leg has two elbows, so a reachable pose has up to four assemblies: out–out, out–in, in–out, in–in. “Out” means the crank–coupler joint is farther from the midline (left P more −x, right Q more +x). This thesis uses out–out, which is the hanging pose drawn in the CAD sketch. Forward kinematics (given θ_L, θ_R) still has two coupler intersections; the hanging robot takes the lower one.",
    )

    para(doc, "4. Singularities to avoid later", bold=True, size=14)
    bullet(doc, "Inverse (leg aligned): |AE| = l1 + l3 (stretched) or |AE| = |l3 − l1| (folded); same for the right leg. The two circle intersections coincide.")
    bullet(doc, "Forward (couplers aligned): P, E, Q collinear. The two forward solutions coincide.")
    para(doc, "None of the five checkpoints is on either singularity. Later trajectories should stay inside the out–out reachable region and away from the workspace rim.")

    para(doc, "5. Checkpoint table (out–out)", bold=True, size=14)
    add_table(
        doc,
        [
            ("Pose", "E (mm)", "θ_L (°)", "θ_R (°)", "link residual", "IK→FK error"),
            ("C1 CAD sample", "(0, −180)", "−125.67", "−54.33", "2.8e-17", "1.1e-16"),
            ("C2 midline high", "(0, −80)", "−138.59", "−41.41", "2.8e-17", "4.9e-15"),
            ("C3 midline low", "(0, −250)", "−104.63", "−75.37", "2.8e-17", "1.4e-17"),
            ("C4 left", "(−60, −180)", "−150.09", "−80.82", "2.8e-17", "2.1e-17"),
            ("C5 right", "(60, −180)", "−99.18", "−29.91", "2.8e-17", "3.1e-17"),
        ],
    )
    para(
        doc,
        "C1 matches the CAD sketch. C4 and C5 are mirrors: θ_L(C4) + θ_R(C5) = −180°. Points (0, −400) mm and (250, 50) mm correctly return no solution. Full numbers are in ik_checkpoints.csv.",
    )

    fig = Path("/workspace/docs/week1-3/figures")
    para(doc, "6. Figures", bold=True, size=14)
    para(doc, "Figure 1. Five checkpoint poses used to verify the program.", sa=4)
    doc.add_picture(str(fig / "five_bar_ik_checkpoints.png"), width=Inches(6.3))
    para(doc, "Figure 2. Four elbow assemblies at the CAD pose. Only out–out is the working mode.", sa=4)
    doc.add_picture(str(fig / "five_bar_ik_four_modes.png"), width=Inches(6.0))
    para(doc, "Figure 3. Reachable workspace of the out–out assembly, with the five checkpoints.", sa=4)
    doc.add_picture(str(fig / "five_bar_ik_workspace.png"), width=Inches(5.2))

    para(doc, "7. How to run", bold=True, size=14)
    para(doc, "python3 scripts/test_five_bar_ik.py")
    para(doc, "python3 scripts/verify_five_bar_ik.py")
    para(
        doc,
        "from five_bar_kinematics import inverse_kinematics",
    )
    para(doc, "sol = inverse_kinematics(0.0, -0.180)   # metres; sol['theta_L_deg'], sol['theta_R_deg']")

    para(doc, "8. Next (still Week 1–3)", bold=True, size=14)
    para(
        doc,
        "Derive the gravitational torques T_w,L and T_w,R from the same geometry and the mid-link COM assumption. Plot T_w along the midline. Do not start GSM / k / ψ until that curve exists.",
    )

    out = Path("/workspace/docs/FiveBar_IK_Week1-3_Bo_Zhang.docx")
    doc.save(out)
    art = Path("/opt/cursor/artifacts")
    if art.is_dir():
        doc.save(art / out.name)
    print("wrote", out)


if __name__ == "__main__":
    build()
