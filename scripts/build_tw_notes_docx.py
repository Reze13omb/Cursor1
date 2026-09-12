#!/usr/bin/env python3
"""Word note for Week 1–3 gravitational torque."""

import csv
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def font(run, size=12, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def para(doc, text, **kw):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(kw.get("sa", 6))
    if kw.get("center"):
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    font(r, size=kw.get("size", 12), bold=kw.get("bold", False), italic=kw.get("italic", False))
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    font(r)


def add_table(doc, rows):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, txt in enumerate(row):
            cell = t.rows[i].cells[j]
            cell.text = ""
            r = cell.paragraphs[0].add_run(str(txt))
            font(r, size=9, bold=(i == 0))
    doc.add_paragraph()


def build():
    csv_path = Path("/workspace/docs/week1-3/tw_checkpoints.csv")
    data = list(csv.DictReader(csv_path.open()))
    by = {r["name"]: r for r in data}

    def row_of(name, label):
        r = by[name]
        return (
            label,
            f"({float(r['x_mm']):.0f}, {float(r['y_mm']):.0f})",
            f"{float(r['T_L_Nm'])*1000:.2f}",
            f"{float(r['T_R_Nm'])*1000:.2f}",
            f"{float(r['F_x_N']):.3f}",
            f"{float(r['F_y_N']):.3f}",
        )

    doc = Document()
    sec = doc.sections[0]
    for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, m, Inches(1))

    para(doc, "Week 1–3 Deliverable 4", bold=True, center=True, sa=2)
    para(doc, "Gravitational Torques of the Hanging Planar Five-Bar", size=16, bold=True, center=True, sa=4)
    para(doc, "Bo Zhang (8571260)  |  Dr Chin-Hsing Kuo  |  September 2026", center=True, sa=10)

    para(doc, "1. What was done", bold=True, size=14)
    para(
        doc,
        "The unbalanced gravitational loads were derived for the frozen geometry and mid-link COM assumption. "
        "Because the five-bar is a closed chain, each actuator torque is not “this leg’s own weight only”. "
        "The correct statement is virtual work: T_w = ∂V/∂θ. "
        "If the cranks are left free, the force F at E that holds the mechanism satisfies J^T F = T_w. "
        "That F is what the IMADA PS-10N would read in a slow hold, before any GSM is fitted. "
        "No springs, gears, or sensors are used this week.",
    )

    para(doc, "2. Potential and holding torque", bold=True, size=14)
    para(doc, "With +y upward and datum on the ground line:", sa=4)
    para(doc, "V = g ( m_a y_aL + m_a y_aR + m_b y_bL + m_b y_bR + m_E y_E )", italic=True)
    para(doc, "Static motor torques that hold the pose:  T_w = (T_wL, T_wR) = ∂V/∂(θ_L, θ_R).", italic=True)
    para(
        doc,
        "COM locations: crank COM at mid-crank; coupler COM at mid-coupler; payload at E. "
        "Masses from the CAD table: m_a = 40 g, m_b = 50 g, m_E = 300 g. "
        "E(θ) comes from the Week 1–3 forward kinematics. "
        "∂E/∂θ is the 2×2 Jacobian Ė = J θ̇ obtained by differentiating the two coupler-length constraints.",
    )

    para(doc, "3. Force at E (matches the later mechanical gauge)", bold=True, size=14)
    para(
        doc,
        "If the two cranks are free (no motor torque), equilibrium is J^T F = T_w. "
        "F = (F_x, F_y) is the external force that must be applied at E. "
        "On the midline, F_x = 0 and F_y is upward (positive), about 3.6–3.7 N for the present masses — inside the 10 N PS-10N range. "
        "This is much smaller than “lift the entire 0.48 kg as a free mass” (4.7 N), because part of the link weight is carried by the grounded joints A and B.",
    )

    para(doc, "4. Checkpoint values", bold=True, size=14)
    add_table(
        doc,
        [
            ("Pose", "E (mm)", "T_wL (mN·m)", "T_wR (mN·m)", "F_x (N)", "F_y (N)"),
            row_of("C1 CAD sample", "C1 CAD sample"),
            row_of("C2 midline high", "C2 midline high (safe)"),
            row_of("C3 midline low", "C3 midline low"),
            row_of("C4 left", "C4 left"),
            row_of("C5 right", "C5 right"),
        ],
    )
    para(
        doc,
        "On the midline, T_wL = −T_wR and F_x = 0 (mirror). "
        "C4 and C5 swap and flip the torques: T_wL(C4) = −T_wR(C5). "
        "Analytic T_w matches a finite-difference of V(θ) to about 10^{-8}–10^{-9} N·m on these poses.",
    )

    para(doc, "5. Singularity to stay away from", bold=True, size=14)
    para(
        doc,
        "Near y ≈ −80 mm on the midline the two couplers line up (forward / Type-II singularity). "
        "The Jacobian is ill-conditioned and the motor torques blow up, even though F_y stays finite. "
        "The IK checkpoint at (0, −80) mm is therefore not used for Tw design. "
        "C2 is taken at (0, −120) mm instead. Later trajectories should stay below about y = −100 mm.",
    )

    fig = Path("/workspace/docs/week1-3/figures/five_bar_tw_midline.png")
    para(doc, "6. Midline curves", bold=True, size=14, sa=4)
    para(doc, "Figure 1. Unbalanced T_w and the hold force at E along x = 0 (y from −260 to −100 mm).", sa=4)
    doc.add_picture(str(fig), width=Inches(5.8))

    para(doc, "7. How to run", bold=True, size=14)
    para(doc, "python3 scripts/test_five_bar_tw.py")
    para(doc, "python3 scripts/verify_five_bar_tw.py")
    para(doc, "from five_bar_statics import gravity_loads")
    para(doc, "data = gravity_loads(0.0, -0.180)  # data['T_L'], data['T_R'], data['F_y']")

    para(doc, "8. Week 1–3 is now complete", bold=True, size=14)
    bullet(doc, "Literature notes")
    bullet(doc, "CAD sketch and frozen dimensions")
    bullet(doc, "Inverse kinematics, four assemblies, workspace")
    bullet(doc, "Unbalanced T_w and the force-gauge prediction F at E")
    para(
        doc,
        "Week 4–6 starts GSM: spring-torque model, targeted configuration, k and ψ. "
        "Do not change d, l1, l3, m_E unless a later print requires a uniform scale.",
    )

    out = Path("/workspace/docs/FiveBar_Tw_Week1-3_Bo_Zhang.docx")
    doc.save(out)
    art = Path("/opt/cursor/artifacts")
    if art.is_dir():
        doc.save(art / out.name)
    print("wrote", out)


if __name__ == "__main__":
    build()
